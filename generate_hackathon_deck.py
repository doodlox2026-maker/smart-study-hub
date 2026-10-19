import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette: Cyberpunk Dark Mode matching web app
    C_BG = RGBColor(15, 23, 42)          # #0F172A Deep Navy
    C_SURFACE = RGBColor(30, 41, 59)     # #1E293B Elevated Slate
    C_CARD_ALT = RGBColor(22, 32, 50)    # #162032 Dark Card
    C_CYAN = RGBColor(14, 165, 233)      # #0EA5E9 Vibrant Cyan Accent
    C_EMERALD = RGBColor(16, 185, 129)   # #10B981 Emerald Green Accent
    C_AMBER = RGBColor(245, 158, 11)     # #F59E0B Gold/Amber Accent
    C_PURPLE = RGBColor(168, 85, 247)    # #A855F7 Purple Accent
    C_ROSE = RGBColor(244, 63, 94)       # #F43F5E Rose Accent
    C_WHITE = RGBColor(248, 250, 252)    # #F8FAFC Pure Crisp White
    C_MUTED = RGBColor(148, 163, 184)    # #94A3B8 Muted Slate

    def add_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = C_BG
        bg_shape.line.fill.background()
        return bg_shape

    def set_slide_transition(slide, t_type="push"):
        try:
            if t_type == "push":
                xml = f'<p:transition {nsdecls("p")} spd="med"><p:push dir="r"/></p:transition>'
            elif t_type == "fade":
                xml = f'<p:transition {nsdecls("p")} spd="med"><p:fade/></p:transition>'
            elif t_type == "wipe":
                xml = f'<p:transition {nsdecls("p")} spd="med"><p:wipe dir="r"/></p:transition>'
            elif t_type == "split":
                xml = f'<p:transition {nsdecls("p")} spd="med"><p:split/></p:transition>'
            elif t_type == "zoom":
                xml = f'<p:transition {nsdecls("p")} spd="med"><p:zoom/></p:transition>'
            else:
                xml = f'<p:transition {nsdecls("p")} spd="med"><p:fade/></p:transition>'
            slide._element.append(parse_xml(xml))
        except Exception as e:
            print(f"Warning setting transition: {e}")

    def add_header(slide, eyebrow, title, subtitle):
        # Eyebrow tag
        eb_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.35))
        tf_eb = eb_box.text_frame
        tf_eb.word_wrap = True
        p_eb = tf_eb.paragraphs[0]
        p_eb.text = eyebrow.upper()
        p_eb.font.size = Pt(10)
        p_eb.font.bold = True
        p_eb.font.color.rgb = C_CYAN

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.65))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE

        # Subtitle
        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.5), Inches(0.4))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.size = Pt(12)
            p_sub.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (HACKATHON SHOWCASE)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_background(s1)
    set_slide_transition(s1, "fade")

    # Glowing decorative top bar
    bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(0.06))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = C_CYAN
    bar1.line.fill.background()

    # Hackathon Badge
    badge1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(3.6), Inches(0.45))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = C_SURFACE
    badge1.line.color.rgb = C_AMBER
    badge1.line.width = Pt(1.5)
    tf_b1 = badge1.text_frame
    p_b1 = tf_b1.paragraphs[0]
    p_b1.text = "🏆 HACKATHON 2026 OFFICIAL SHOWCASE"
    p_b1.font.size = Pt(11)
    p_b1.font.bold = True
    p_b1.font.color.rgb = C_AMBER
    p_b1.alignment = PP_ALIGN.CENTER

    # Main Project Title
    t1_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.95), Inches(11.733), Inches(1.5))
    tf1 = t1_box.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "SMART STUDY HUB"
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Next-Gen Interactive Academic Concept Studio & Mind Map Engine"
    p1_sub.font.size = Pt(20)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = C_CYAN

    # Elevator Pitch Box
    pitch_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.8), Inches(11.733), Inches(1.5))
    pitch_card.fill.solid()
    pitch_card.fill.fore_color.rgb = C_SURFACE
    pitch_card.line.color.rgb = C_CYAN
    pitch_card.line.width = Pt(1)
    tf_pitch = pitch_card.text_frame
    tf_pitch.margin_left = Inches(0.3)
    tf_pitch.margin_right = Inches(0.3)
    p_pitch = tf_pitch.paragraphs[0]
    p_pitch.text = "💡 ELEVATOR PITCH:"
    p_pitch.font.size = Pt(12)
    p_pitch.font.bold = True
    p_pitch.font.color.rgb = C_EMERALD

    p_pitch_body = tf_pitch.add_paragraph()
    p_pitch_body.text = (
        "A zero-dependency, ultra-fast academic workspace engineered to eliminate textbook fatigue. "
        "Fuses rigorous STEM and Humanities theory with interactive SVG vector simulations, massive per-subject "
        "concept mind maps, a live focus stopwatch, and an evaluation engine with instant positive reinforcement and deep error explanations."
    )
    p_pitch_body.font.size = Pt(13)
    p_pitch_body.font.color.rgb = C_WHITE

    # Key Highlights Badges at bottom
    badges_data = [
        ("⚡ 100% Offline Ready", C_CYAN),
        ("🎯 40 Curated Questions", C_EMERALD),
        ("🧠 Huge Mind Maps in 4 Subjects", C_PURPLE),
        ("🚀 Zero External Dependencies", C_AMBER)
    ]
    for idx, (label, col) in enumerate(badges_data):
        bx = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 3.0), Inches(5.6), Inches(2.8), Inches(0.6))
        bx.fill.solid()
        bx.fill.fore_color.rgb = C_CARD_ALT
        bx.line.color.rgb = col
        bx.line.width = Pt(1)
        tf_bx = bx.text_frame
        p_bx = tf_bx.paragraphs[0]
        p_bx.text = label
        p_bx.font.size = Pt(11)
        p_bx.font.bold = True
        p_bx.font.color.rgb = col
        p_bx.alignment = PP_ALIGN.CENTER

    # Footer Info
    foot_box = s1.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.4))
    tf_foot = foot_box.text_frame
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "Track: EdTech & Intelligent Systems  •  Live Deployment: GitHub Pages  •  Team: Doodlox2026-Maker"
    p_foot.font.size = Pt(11)
    p_foot.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 2: THE PROBLEM (THE PAIN POINT)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_background(s2)
    set_slide_transition(s2, "push")
    add_header(s2, "The Student Dilemma", "Why Modern Academic Revision is Broken", "Identifying the root causes of cognitive overload and poor retention in academic learning.")

    problems = [
        ("1. Abstract Formulas Lack Visual Intuition", 
         "Students memorize calculus limits, ethyne orbital bonds, and electromagnetic flux equations as dry symbols without seeing how parameters dynamically move or behave in physical reality.",
         C_ROSE),
        ("2. Disjointed Notes & Missing Topology", 
         "Standard textbooks treat topics as isolated silos. Learners fail to visualize how foundational axioms in algebra or kinematics bridge into advanced physical and computational systems.",
         C_AMBER),
        ("3. Quizzes That Say 'Wrong' Without Teaching", 
         "Traditional test tools simply mark options red or green. They offer zero academic breakdown explaining why a misconception is incorrect or how the right principle applies.",
         C_PURPLE)
    ]

    for i, (title, desc, accent) in enumerate(problems):
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.95), Inches(2.1), Inches(3.75), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_SURFACE
        card.line.color.rgb = accent
        card.line.width = Pt(1.5)
        tf = card.text_frame
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.3)
        
        p_icon = tf.paragraphs[0]
        p_icon.text = f"🚨 GAP {i+1}"
        p_icon.font.size = Pt(11)
        p_icon.font.bold = True
        p_icon.font.color.rgb = accent
        
        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE
        p_t.space_before = Pt(8)
        p_t.space_after = Pt(14)
        
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 3: THE SOLUTION (SMART STUDY HUB)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_background(s3)
    set_slide_transition(s3, "push")
    add_header(s3, "The Breakthrough Solution", "Smart Study Hub: Full-Spectrum Learning Architecture", "An integrated, zero-lag ecosystem delivering rigorous conceptual depth and active testing.")

    pillars = [
        ("📐 4 Core Disciplines", "Mathematics (Calculus/Matrices), Chemistry (Organic/Bonding), Physics (Mechanics/Induction), and English (Grammar/Literature) nested systematically.", C_CYAN),
        ("🎨 Vector Simulation Studio", "Interactive SVG vector models rendering tangent lines, ethyne molecular pi-clouds, velocity-time acceleration/displacement integrals, and syntax trees.", C_EMERALD),
        ("🧠 Huge Concept Mind Maps", "Every subject possesses an expansive dedicated knowledge graph connecting foundational theorems to applied engineering and critical analysis.", C_PURPLE),
        ("🎯 40-Question Evaluation", "10 curated questions per subject with cheering victory complements for correct answers and deep academic error explanations for mistakes.", C_AMBER),
        ("⭕ Live Progress Wheel", "Hardware-accelerated SVG circular progress wheel that continuously calculates and animates percentage accuracy and question breakdown in real-time.", C_CYAN),
        ("⏱️ Precision Focus Stopwatch", "Active session stopwatch counting up second-by-second to track study blocks with pause/resume controls and strict exit confirmation.", C_EMERALD)
    ]

    for idx, (title, text, col) in enumerate(pillars):
        row = idx // 3
        col_idx = idx % 3
        cx = Inches(0.8 + col_idx * 3.95)
        cy = Inches(2.1 + row * 2.35)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(3.75), Inches(2.1))
        card.fill.solid()
        card.fill.fore_color.rgb = C_SURFACE
        card.line.color.rgb = col
        card.line.width = Pt(1)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)
        
        p_h = tf.paragraphs[0]
        p_h.text = title
        p_h.font.size = Pt(14)
        p_h.font.bold = True
        p_h.font.color.rgb = col
        
        p_b = tf.add_paragraph()
        p_b.text = text
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = C_MUTED
        p_b.space_before = Pt(6)

    # =========================================================================
    # SLIDE 4: INTERACTIVE VISUAL CONCEPT STUDIO (THEORY LAB)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_background(s4)
    set_slide_transition(s4, "wipe")
    add_header(s4, "Visual Knowledge Systems", "Interactive Visual Concept Studio & Vector Models", "Translating complex mathematical and scientific equations into intuitive vector animations.")

    visual_models = [
        ("📐 Mathematics", "Calculus Tangents & Matrix Area", 
         "• Cauchy-Weierstrass limits & secant convergence\n• Tangent slope m = f'(x) at point P\n• Determinant |ad - bc| as 2D area scale factor\n• Singularity collapses vector dimensionality", C_CYAN),
        ("🧪 Chemistry", "Orbital Hybridization & Ionic Lattice", 
         "• Ethyne (C2H2) linear 180° sp geometry\n• Central σ-core with orthogonal π-electron clouds\n• Na⁺ and Cl⁻ complete valence electron transfer\n• Coulombic crystal lattice attraction", C_EMERALD),
        ("⚛️ Physics", "Velocity-Time Calculus & Faraday Coil", 
         "• v-t graph instantaneous slope = acceleration a\n• Definite integral ∫ v dt = displacement area s\n• Solenoid magnetic induction opposing flux ingress\n• Lenz's Law thermodynamic energy conservation", C_AMBER),
        ("📖 English", "Syntax Parse Trees & Scansion Waves", 
         "• Complex sentence dependency trees\n• Subordinate vs independent clause hierarchy\n• Shakespeare Sonnet 18 Iambic Pentameter\n• Acoustic stress waveforms across 5 metric feet", C_PURPLE)
    ]

    for idx, (sub, title, details, col) in enumerate(visual_models):
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 2.95), Inches(2.1), Inches(2.8), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_SURFACE
        card.line.color.rgb = col
        card.line.width = Pt(1.5)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.25)
        
        p1 = tf.paragraphs[0]
        p1.text = sub
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = col
        
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE
        p2.space_before = Pt(4)
        p2.space_after = Pt(10)
        
        p3 = tf.add_paragraph()
        p3.text = details
        p3.font.size = Pt(11)
        p3.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 5: HUGE MIND MAP STUDIO (IN EVERY SUBJECT)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_background(s5)
    set_slide_transition(s5, "push")
    add_header(s5, "Knowledge Topologies", "Huge Mind Map Studio (Per-Subject Deep Graphs)", "Expansive visual trees mapping foundational axioms into cross-disciplinary applications.")

    maps_data = [
        ("📐 Mathematics Mind Map", "Central Root: Mathematics Core Topology", 
         "Branch 1: Calculus & Mathematical Analysis\n  → Limits & Continuity (Squeeze, L'Hôpital)\n  → Derivatives & Tangent Rates\n  → Power, Product & Chain Rules\n  → Optimization & Critical Points\n\nBranch 2: Linear Algebra & Matrix Systems\n  → Matrix Operations (Dot products)\n  → Determinants (ad - bc volume scaling)\n  → Matrix Invertibility & Singularity\n  → Linear Systems & Cramer's Rule"),
        ("🧪 Chemistry Mind Map", "Central Root: Chemical Molecular Topology", 
         "Branch 1: Organic Chemistry & Hydrocarbons\n  → Alkanes, Alkenes & Alkynes\n  → sp³ / sp² / sp Hybridization Geometries\n  → Aromatic Benzene & Hückel (4n+2)\n  → Functional Families (Alcohols, Acids)\n\nBranch 2: Periodic Table & Chemical Bonding\n  → Periodic Electronegativity Trends\n  → Ionic Crystal Lattices (Coulomb)\n  → Covalent Electron Sharing & Octets\n  → Intermolecular Hydrogen Bonding"),
        ("⚛️ Physics Mind Map", "Central Root: Physics Dynamics Topology", 
         "Branch 1: Classical Mechanics & Motion\n  → Rectilinear Kinematic Equations\n  → V-T Graph Calculus (Slope & Area)\n  → Free Fall & Gravitational Trajectory\n  → Work-Energy Invariant v² = u² + 2as\n\nBranch 2: Electromagnetism & Field Theory\n  → Magnetic Flux Surface Integrals (Webers)\n  → Faraday's Law of Induction\n  → Lenz's Law Thermodynamic Opposition\n  → Transformers & Eddy Current Mitigation"),
        ("📖 English Mind Map", "Central Root: Linguistic & Poetic Topology", 
         "Branch 1: Grammar & Syntactic Structures\n  → Clause Dependency Hierarchy\n  → Tense Aspect Matrix (Past Perfect)\n  → Subject-Verb Agreement & Concord\n  → Semicolons & Run-On Sentence Repairs\n\nBranch 2: Literature & Poetic Criticism\n  → Iambic Pentameter Scansion (5 feet)\n  → Conceits & Extended Metaphors\n  → Shakespearean Sonnet & Volta\n  → Phonetic Alliteration & Personification")
    ]

    for idx, (title, root, branches) in enumerate(maps_data):
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 2.95), Inches(2.1), Inches(2.8), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = C_SURFACE
        card.line.color.rgb = C_CYAN if idx % 2 == 0 else C_EMERALD
        card.line.width = Pt(1)
        tf = card.text_frame
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)
        tf.margin_top = Inches(0.2)
        
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE
        
        p_r = tf.add_paragraph()
        p_r.text = root
        p_r.font.size = Pt(10)
        p_r.font.bold = True
        p_r.font.color.rgb = C_AMBER
        p_r.space_before = Pt(4)
        p_r.space_after = Pt(8)
        
        p_b = tf.add_paragraph()
        p_b.text = branches
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 6: EVALUATION ENGINE & PROGRESS ACCURACY WHEEL
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_background(s6)
    set_slide_transition(s6, "push")
    add_header(s6, "Evaluation Engine", "Curated 40-Question System & Dynamic Accuracy Wheel", "Active retrieval testing backed by immediate positive reinforcement and explainable feedback.")

    # Left: 40-Question Engine Card
    q_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.1), Inches(6.8), Inches(4.3))
    q_card.fill.solid()
    q_card.fill.fore_color.rgb = C_SURFACE
    q_card.line.color.rgb = C_EMERALD
    q_card.line.width = Pt(1.5)
    tf_q = q_card.text_frame
    tf_q.margin_left = Inches(0.3)
    tf_q.margin_right = Inches(0.3)
    tf_q.margin_top = Inches(0.3)
    
    p_qh = tf_q.paragraphs[0]
    p_qh.text = "🎯 40 Curated Academic Benchmarks (10 Per Subject)"
    p_qh.font.size = Pt(16)
    p_qh.font.bold = True
    p_qh.font.color.rgb = C_WHITE
    
    p_qb = tf_q.add_paragraph()
    p_qb.text = (
        "• Mathematics (10 Qs): Limits, Determinants, Power/Product Rules, Singularity, Tangents, Euler's e.\n"
        "• Chemistry (10 Qs): Alkynes, Ionic Lattices, Alcohols, Electronegativity, Benzene, Octet Rule.\n"
        "• Physics (10 Qs): Kinematics Displacement, Lenz's Law, v-t Slopes, Webers, Transformers, Free Fall.\n"
        "• English (10 Qs): Complex Clauses, Sonnet Conceits, Subject-Verb Concord, Iambic Pentameter.\n\n"
        "✨ Positive Reinforcement:\n"
        "Correct answers trigger green success compliments cheering the student on.\n\n"
        "💡 Deep Explainable Feedback:\n"
        "Wrong answers trigger thorough academic analyses directly explaining WHY that selection was flawed and proving WHY the correct answer holds true."
    )
    p_qb.font.size = Pt(11.5)
    p_qb.font.color.rgb = C_MUTED
    p_qb.space_before = Pt(8)

    # Right: Circular Progress Wheel Card
    w_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(2.1), Inches(4.6), Inches(4.3))
    w_card.fill.solid()
    w_card.fill.fore_color.rgb = C_CARD_ALT
    w_card.line.color.rgb = C_CYAN
    w_card.line.width = Pt(1.5)
    tf_w = w_card.text_frame
    tf_w.margin_left = Inches(0.3)
    tf_w.margin_right = Inches(0.3)
    tf_w.margin_top = Inches(0.3)
    
    p_wh = tf_w.paragraphs[0]
    p_wh.text = "⭕ Real-Time Accuracy Progress Wheel"
    p_wh.font.size = Pt(15)
    p_wh.font.bold = True
    p_wh.font.color.rgb = C_CYAN
    
    p_wb = tf_w.add_paragraph()
    p_wb.text = (
        "• Hardware-Accelerated Vector SVG Arc:\n"
        "  Animates stroke-dashoffset transitions smoothly based on exact real-time performance.\n\n"
        "• Dynamic Numerical HUD:\n"
        "  Displays real-time percentage accuracy (0% - 100%).\n\n"
        "• Instant Performance Metric Breakdown:\n"
        "  - Correct Answers (Emerald indicator)\n"
        "  - Incorrect Answers (Rose indicator)\n"
        "  - Remaining Pool (Slate indicator)\n"
        "  - Academic Standing (Standby, Proficient, High Distinction)"
    )
    p_wb.font.size = Pt(11.5)
    p_wb.font.color.rgb = C_MUTED
    p_wb.space_before = Pt(8)

    # =========================================================================
    # SLIDE 7: TECH STACK & COMPETITIVE ADVANTAGE
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_background(s7)
    set_slide_transition(s7, "wipe")
    add_header(s7, "Engineering Excellence", "Zero-Dependency Architecture vs Legacy Bloat", "Maximizing speed, accessibility, and offline durability through pure web standards.")

    stack_cards = [
        ("⚡ 100% Zero Dependencies", "No React, Vue, Angular, or external UI CDNs. Entire application compiles into a single, self-contained index.html file that loads in <50ms.", C_EMERALD),
        ("🎨 Hardware-Accelerated SVG", "All mathematical curves, chemical orbital clouds, physics kinematic graphs, and clause trees are native vector SVGs rendered at 60 FPS.", C_CYAN),
        ("⏱️ Precision Event-Loop Engine", "Real-time non-blocking stopwatch and evaluation accuracy calculations execute instantaneously without battery drain or lag.", C_AMBER),
        ("🌐 Edge CDN Static Deployment", "Hosted directly on GitHub Pages with global CDN caching. Zero server maintenance, zero database downtime, and zero cost.", C_PURPLE)
    ]

    for idx, (title, desc, col) in enumerate(stack_cards):
        cx = Inches(0.8 + (idx % 2) * 5.95)
        cy = Inches(2.1 + (idx // 2) * 2.35)
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, Inches(5.75), Inches(2.1))
        card.fill.solid()
        card.fill.fore_color.rgb = C_SURFACE
        card.line.color.rgb = col
        card.line.width = Pt(1)
        tf = card.text_frame
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = col
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = C_MUTED
        p2.space_before = Pt(6)

    # =========================================================================
    # SLIDE 8: LIVE DEMO & JUDGE TESTING
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_background(s8)
    set_slide_transition(s8, "push")
    add_header(s8, "Live Verification", "Test the Live Deployment Right Now", "Open and explore Smart Study Hub on any mobile phone, tablet, or laptop immediately.")

    # Live Link Banner
    link_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.1), Inches(11.733), Inches(1.3))
    link_card.fill.solid()
    link_card.fill.fore_color.rgb = C_CARD_ALT
    link_card.line.color.rgb = C_CYAN
    link_card.line.width = Pt(1.5)
    tf_link = link_card.text_frame
    p_lh = tf_link.paragraphs[0]
    p_lh.text = "🌐 PUBLIC LIVE GITHUB PAGES URL:"
    p_lh.font.size = Pt(12)
    p_lh.font.bold = True
    p_lh.font.color.rgb = C_CYAN
    
    p_lurl = tf_link.add_paragraph()
    p_lurl.text = "https://doodlox2026-maker.github.io/smart-study-hub/smart-study-hub/"
    p_lurl.font.size = Pt(20)
    p_lurl.font.bold = True
    p_lurl.font.color.rgb = C_WHITE
    p_lurl.space_before = Pt(4)

    # 3 Steps to Demo
    demo_steps = [
        ("Step 1: Open on Mobile or Desktop", "No account, login, or install required. Loads instantly in any standard browser.", C_CYAN),
        ("Step 2: Touch Any Subject", "Inspect the Visual Studio simulations and explore the Huge Dedicated Mind Map.", C_EMERALD),
        ("Step 3: Test the 40 Questions", "Answer questions across Math, Chem, Physics & English to test real-time feedback & accuracy wheel.", C_AMBER)
    ]
    for idx, (title, desc, col) in enumerate(demo_steps):
        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 3.95), Inches(3.7), Inches(3.75), Inches(2.7))
        card.fill.solid()
        card.fill.fore_color.rgb = C_SURFACE
        card.line.color.rgb = col
        card.line.width = Pt(1)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)
        
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = col
        
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = C_MUTED
        p_d.space_before = Pt(8)

    # =========================================================================
    # SLIDE 9: FUTURE ROADMAP & THANK YOU
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_background(s9)
    set_slide_transition(s9, "fade")
    add_header(s9, "Conclusion & Impact", "Thank You! Questions & Discussion", "The future of interactive, accessible, and high-yield academic learning.")

    # Left: Roadmap
    r_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.1), Inches(6.5), Inches(4.3))
    r_card.fill.solid()
    r_card.fill.fore_color.rgb = C_SURFACE
    r_card.line.color.rgb = C_PURPLE
    r_card.line.width = Pt(1.5)
    tf_r = r_card.text_frame
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_right = Inches(0.3)
    tf_r.margin_top = Inches(0.3)
    
    p_rh = tf_r.paragraphs[0]
    p_rh.text = "🔮 Future Engineering Roadmap"
    p_rh.font.size = Pt(16)
    p_rh.font.bold = True
    p_rh.font.color.rgb = C_PURPLE
    
    p_rb = tf_r.add_paragraph()
    p_rb.text = (
        "• Spaced Repetition Flashcards (SuperMemo SM-2 / Anki Algorithm):\n"
        "  Automating long-term retention schedules for difficult conceptual topics.\n\n"
        "• Voice-Guided Formula Synthesis:\n"
        "  Audio narration explaining mathematical and chemical derivations.\n\n"
        "• PDF Performance Dossier Generator:\n"
        "  One-click export of student revision analytics, time spent, and mastery ratings.\n\n"
        "• Expanded Curriculum Tracks:\n"
        "  Biology, Computer Science & Algorithms, and Modern History."
    )
    p_rb.font.size = Pt(12)
    p_rb.font.color.rgb = C_MUTED
    p_rb.space_before = Pt(8)

    # Right: Thank You & Contact
    ty_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.6), Inches(2.1), Inches(4.933), Inches(4.3))
    ty_card.fill.solid()
    ty_card.fill.fore_color.rgb = C_CARD_ALT
    ty_card.line.color.rgb = C_CYAN
    ty_card.line.width = Pt(1.5)
    tf_ty = ty_card.text_frame
    tf_ty.margin_left = Inches(0.3)
    tf_ty.margin_right = Inches(0.3)
    tf_ty.margin_top = Inches(0.4)
    
    p_tyh = tf_ty.paragraphs[0]
    p_tyh.text = "🏆 SMART STUDY HUB"
    p_tyh.font.size = Pt(22)
    p_tyh.font.bold = True
    p_tyh.font.color.rgb = C_WHITE
    
    p_tyb = tf_ty.add_paragraph()
    p_tyb.text = (
        "Empowering students through visual intuition, deep theory, and instant evaluation.\n\n"
        "💻 GitHub Repository:\n"
        "github.com/doodlox2026-maker/smart-study-hub\n\n"
        "🌐 Live Web App:\n"
        "doodlox2026-maker.github.io/smart-study-hub/smart-study-hub/\n\n"
        "Ready for Judge Q&A! 🚀"
    )
    p_tyb.font.size = Pt(13)
    p_tyb.font.color.rgb = C_MUTED
    p_tyb.space_before = Pt(10)

    # Save to Downloads and smart-study-hub
    out_path_1 = r"c:\Users\MARI KUMAR\Downloads\smart_study_hub_hackathon_pitch.pptx"
    out_path_2 = r"c:\Users\MARI KUMAR\Downloads\smart-study-hub\smart_study_hub_hackathon_pitch.pptx"

    prs.save(out_path_1)
    prs.save(out_path_2)
    print(f"Successfully generated: {out_path_1}")
    print(f"Successfully generated: {out_path_2}")

if __name__ == "__main__":
    create_deck()
